"""An .xlsx whose formulas use defined names, for the Sheets names tests (R148).

Names of each kind Excel files carry: a range with a comment (Revenue), a
cell (Rate), a value (TaxRate), a reference to a sheet whose name needs
quotes (Other), one over two areas (Two, which Sheets leaves out), and one
scoped to a sheet (Local). Excel's saved answers are written into each
formula cell, as Excel would save them (openpyxl writes none), so the import
can be checked to compute rather than show them.

    python tests/fixtures/sheets/make_openpyxl_names.py
"""

import os
import re
import zipfile

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-names.xlsx")
TMP = OUT + ".tmp"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"
ws["B1"], ws["D1"], ws["E1"] = "Revenue", "Rate", 0.1
for r, v in ((2, 10), (3, 20), (4, 30)):
    ws.cell(r, 2, v)
q = wb.create_sheet("Q 3")
q["A1"] = 5

# Each formula and the answer Excel gives it.
FORMULAS = {
    ("Data", "D3"): ("=SUM(Revenue)", 60),
    ("Data", "D4"): ("=SUM(Revenue)*Rate", 6),
    ("Data", "D5"): ("=SUM(B2:B4)", 60),
    ("Data", "D6"): ("=SUM(Revenue)*TaxRate", 12),
    ("Data", "D7"): ("=Other*2", 10),
    ("Data", "D8"): ("=Local*100", 10),
    ("Data", "D9"): ("=SUM(Two)", 40),
    ("Q 3", "B1"): ("=SUM(Revenue)", 60),
    ("Q 3", "B2"): ("=ROWS(Revenue)", 3),
}
for (sheet, addr), (f, _) in FORMULAS.items():
    wb[sheet][addr] = f

wb.defined_names["Revenue"] = DefinedName(
    "Revenue", attr_text="Data!$B$2:$B$4", comment="Monthly revenue"
)
wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$E$1")
wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="0.2")
wb.defined_names["Other"] = DefinedName("Other", attr_text="'Q 3'!$A$1")
wb.defined_names["Two"] = DefinedName("Two", attr_text="Data!$B$2,Data!$B$4")
ws.defined_names["Local"] = DefinedName("Local", attr_text="Data!$E$1")
wb.save(TMP)

PARTS = {"Data": "xl/worksheets/sheet1.xml", "Q 3": "xl/worksheets/sheet2.xml"}
with zipfile.ZipFile(TMP) as zin, zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        for sheet, part in PARTS.items():
            if item.filename != part:
                continue
            xml = data.decode("utf-8")
            for (s, addr), (_, v) in FORMULAS.items():
                if s != sheet:
                    continue
                pat = re.compile(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)(<v\s*/>|<v>[^<]*</v>)?' % addr)
                xml, n = pat.subn(lambda m: m.group(1) + m.group(2) + "<v>%s</v>" % v, xml)
                assert n == 1, (sheet, addr, n)
            data = xml.encode("utf-8")
        zout.writestr(item, data)
os.remove(TMP)
print(OUT)
