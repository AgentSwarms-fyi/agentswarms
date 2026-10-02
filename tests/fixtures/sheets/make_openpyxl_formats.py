"""Number formats Excel shows in ways a quick reading misses (R163).

Each row: a value, its format, and what Excel shows (column C, text).
Excel's own built-in format 46 is [h]:mm:ss, so every duration in a
timesheet comes this way; [$€-2] is how a euro format is saved; fractions
and conditions are Excel's own Fraction and Special categories.

    python tests/fixtures/sheets/make_openpyxl_formats.py
"""

import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-formats.xlsx")

ROWS = [
    (1.5, "[h]:mm", "36:00"),
    (1.5, "[h]:mm:ss", "36:00:00"),
    (0.0625, "[mm]:ss", "90:00"),
    (0.0625, "[ss]", "5400"),
    (2.75, "[h]\" hours\"", "66 hours"),
    (1234.5, "[$€-2] #,##0.00", "€ 1,234.50"),
    (1234.5, "[$£-809]#,##0.00", "£1,234.50"),
    (-1234.5, "[$$-409]#,##0.00;[Red]-[$$-409]#,##0.00", "-$1,234.50"),
    (1.5, "# ?/?", "1 1/2"),
    (0.75, "?/?", "3/4"),
    (2.625, "# ?/8", "2 5/8"),
    (12, "[<10]\"small\";\"big\"", "big"),
    (5, "[<10]\"small\";\"big\"", "small"),
    (150, "[>=100]\"high\";[<50]\"low\";\"mid\"", "high"),
    (75, "[>=100]\"high\";[<50]\"low\";\"mid\"", "mid"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Formats"
ws.append(["Value", "Format", "Excel shows"])
for value, code, shown in ROWS:
    ws.append([value, code, shown])
    ws.cell(ws.max_row, 1).number_format = code
ws.column_dimensions["B"].width = 44
ws.column_dimensions["C"].width = 16
wb.save(OUT)
print(OUT)
