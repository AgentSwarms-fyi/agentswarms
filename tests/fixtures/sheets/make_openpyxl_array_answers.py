"""Functions that answer with several values, written as older Excel wrote them (R338).

Typed in one cell without Ctrl+Shift+Enter, such a formula showed one value
in the Excel that wrote it, and Excel 365 shows an @ in front of the
function: Microsoft's page on @ gives =@INDEX(A1:A10,B1) and
=@OFFSET(A1:A2,1,1), except where the function is "wrapped in a function
that accepts an array or range (for example, SUM() or AVERAGE())". A range
it answers with is read in the formula's own row; an array gives its first
value. openpyxl writes every formula plain, as older Excel did. The expected
values are older Excel's, worked out by hand:

  Legacy   A Month  B Sales
           1        3100
           2        4500
           3        4400
           4        5400
           5        7500
           6        8100      (slope 1000, intercept 2000)

  D2   =LINEST(B2:B7,A2:A7)                   the first value, the slope: 1000
  D3   =SUM(LINEST(B2:B7,A2:A7)*{9,1})        SUM takes the array (LINEST's page): 11000
  D4   =INDEX(LINEST(B2:B7,A2:A7),2)          the intercept: 2000
  D5   =ROW(A2:A4)                            the first row: 2
  D6   =SUMPRODUCT((B2:B7>5000)*ROW(B2:B7))   rows 5, 6 and 7: 18
  D7   =TRANSPOSE(A2:A4)                      the first value: 1
  D8   =INDEX(A2:B7,0,2)                      B2:B7 in row 8, outside it: #VALUE!
  D9   =ABS(-LINEST(B2:B7,A2:A7))             1000
  D10  =TREND(B2:B7,A2:A7,{7;8})              month 7: 9000
  D12  {=LINEST(B2:B7,A2:A7)}                 an array formula over D12:E12: 1000 and 2000
  D13  =ROW()                                 13
  D14  =ROW(A2)                               2
  F3   =OFFSET(A2:A3,1,1)                     B3:B4 in row 3: 4500
  F4   =AVERAGE(OFFSET(A2,0,1,6,1))           AVERAGE takes the range: 5500
  F5   =INDEX(A2:B7,0,2)                      B2:B7 in row 5: 5400
  F6   =IF(A2>0,B2:B7,0)                      the range IF answers with, in row 6: 7500

    python tests/fixtures/sheets/make_openpyxl_array_answers.py
"""

import os

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-array-answers.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Legacy"
ws.append(["Month", "Sales"])
for row in [(1, 3100), (2, 4500), (3, 4400), (4, 5400), (5, 7500), (6, 8100)]:
    ws.append(row)
ws["D2"] = "=LINEST(B2:B7,A2:A7)"
ws["D3"] = "=SUM(LINEST(B2:B7,A2:A7)*{9,1})"
ws["D4"] = "=INDEX(LINEST(B2:B7,A2:A7),2)"
ws["D5"] = "=ROW(A2:A4)"
ws["D6"] = "=SUMPRODUCT((B2:B7>5000)*ROW(B2:B7))"
ws["D7"] = "=TRANSPOSE(A2:A4)"
ws["D8"] = "=INDEX(A2:B7,0,2)"
ws["D9"] = "=ABS(-LINEST(B2:B7,A2:A7))"
ws["D10"] = "=TREND(B2:B7,A2:A7,{7;8})"
ws["D12"] = ArrayFormula("D12:E12", "=LINEST(B2:B7,A2:A7)")
ws["D13"] = "=ROW()"
ws["D14"] = "=ROW(A2)"
ws["F3"] = "=OFFSET(A2:A3,1,1)"
ws["F4"] = "=AVERAGE(OFFSET(A2,0,1,6,1))"
ws["F5"] = "=INDEX(A2:B7,0,2)"
ws["F6"] = "=IF(A2>0,B2:B7,0)"
wb.save(OUT)
print(OUT)
