"""A workbook whose visible sheet reads a hidden and a very hidden sheet (R160).

Excel keeps helper sheets hidden: "Rates" is hidden (Format > Hide), and
"Keys" is very hidden (only VBA or the file can show it). The Summary's
formulas read both, so an import that drops either breaks them.

    python tests/fixtures/sheets/make_openpyxl_hidden.py
"""

import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-hidden.xlsx")

wb = openpyxl.Workbook()
summary = wb.active
summary.title = "Summary"
summary["A1"], summary["B1"] = "Region", "Revenue (EUR)"
for i, (region, usd) in enumerate([("West", 1000), ("East", 2500), ("North", 400)], start=2):
    summary.cell(i, 1, region)
    summary.cell(i, 3, usd)
    summary.cell(i, 2, f"=C{i}*Rates!$B$2")
summary["A6"] = "Total"
summary["B6"] = "=SUM(B2:B4)"
summary["A7"] = "Key"
summary["B7"] = "=Keys!A1"

rates = wb.create_sheet("Rates")
rates["A1"], rates["B1"] = "Currency", "Rate"
rates["A2"], rates["B2"] = "EUR", 0.9
rates.sheet_state = "hidden"

keys = wb.create_sheet("Keys")
keys["A1"] = "K-2026"
keys.sheet_state = "veryHidden"

wb.save(OUT)
print(OUT)
