"""Formulas written the way Excel before dynamic arrays wrote them (R162).

A plain formula in a file is read by Excel 365 as an older one: where it
expects one value and meets a range, it takes the value in the formula's
own row (and shows @ there). openpyxl, like older Excel, writes every
formula plain. The expected values below are older Excel's, worked out by
hand from the rows:

  Orders   A Item   B Price  C Qty
           Pen      2        5
           Ink      10       1
           Pad      4        3

  D2:D4  =Price*Qty        names over whole columns: 10, 10, 12
  E2:E4  =B:B*C:C          the same, with the columns written out
  F2     =A2:A4            the row's own item: Pen (F3 would be Ink)
  G5     =A2:A4            row 5 is outside A2:A4: #VALUE!
  H2     =SUM(LEN(A2:A4))  LEN of the row's item: 3 (not 9)
  I2     =SUMPRODUCT((C2:C4>1)*B2:B4)   SUMPRODUCT works over arrays in every Excel: 2+4 = 6
  J2     =LOOKUP(2,1/(C2:C4>0),B2:B4)   so does LOOKUP: the last price, 4
  K2     =SUM(B2:B4)       a range where Excel always takes one: 16
  L2     {=SUM(LEN(A2:A4))}  an array formula (Ctrl+Shift+Enter): 9
  M2     =_xlfn.SINGLE(A2:A4)&"!"   Excel 365's own @, as its files hold it: Pen!

    python tests/fixtures/sheets/make_openpyxl_legacy.py
"""

import os

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-legacy.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Orders"
ws.append(["Item", "Price", "Qty"])
for row in [("Pen", 2, 5), ("Ink", 10, 1), ("Pad", 4, 3)]:
    ws.append(row)
for r in range(2, 5):
    ws[f"D{r}"] = "=Price*Qty"
    ws[f"E{r}"] = "=B:B*C:C"
ws["F2"] = "=A2:A4"
ws["G5"] = "=A2:A4"
ws["H2"] = "=SUM(LEN(A2:A4))"
ws["I2"] = "=SUMPRODUCT((C2:C4>1)*B2:B4)"
ws["J2"] = "=LOOKUP(2,1/(C2:C4>0),B2:B4)"
ws["K2"] = "=SUM(B2:B4)"
ws["L2"] = ArrayFormula("L2", "=SUM(LEN(A2:A4))")
ws["M2"] = '=_xlfn.SINGLE(A2:A4)&"!"'
wb.defined_names["Price"] = DefinedName("Price", attr_text="Orders!$B:$B")
wb.defined_names["Qty"] = DefinedName("Qty", attr_text="Orders!$C:$C")
wb.save(OUT)
print(OUT)
