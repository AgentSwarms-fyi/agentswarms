"""A contact list with repeated rows, for the Sheets Remove Duplicates tests (R153).

Excel's Data → Remove Duplicates compares what each cell shows, ignoring case:
"ASHA@EXAMPLE.COM" repeats "asha@example.com", and the same date shown two ways
("2026-03-08" and "8 Mar 2026") does not repeat. Column G sits outside the
list, past an empty column, and must not move.

    python tests/fixtures/sheets/make_openpyxl_dupes.py
"""

import datetime as dt
import os

import openpyxl
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-dupes.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Contacts"
ws.append(["Name", "Email", "City", "Signed up", "Chars"])
rows = [
    ("Asha Rao", "asha@example.com", "Lisbon", dt.date(2026, 3, 8)),
    ("Ben Ode", "ben@example.com", "Porto", dt.date(2026, 3, 9)),
    ("Asha Rao", "ASHA@EXAMPLE.COM", "Lisbon", dt.date(2026, 3, 8)),  # a repeat, in capitals
    ("Cleo Park", "cleo@example.com", "Faro", dt.date(2026, 3, 10)),
    ("Ben Ode", "ben@example.com", "Porto", dt.date(2026, 3, 9)),  # a repeat
    ("Dev Iyer", "dev@example.com", "Braga", dt.date(2026, 3, 8)),
    ("Dev Iyer", "dev@example.com", "Braga", dt.date(2026, 3, 8)),  # shown another way below
]
for i, (name, email, city, day) in enumerate(rows, start=2):
    ws.cell(i, 1, name)
    ws.cell(i, 2, email)
    ws.cell(i, 3, city)
    ws.cell(i, 4, day).number_format = "yyyy-mm-dd"
    ws.cell(i, 5, f"=LEN(B{i})")
# The last Dev row's date, shown as "8 Mar 2026": the same day, not the same text.
ws["D8"].number_format = "d mmm yyyy"
ws["C5"].comment = Comment("Moved from Lisbon in April.", "Ops")
for col, width in zip("ABCDE", (14, 22, 10, 13, 8)):
    ws.column_dimensions[col].width = width

# Past an empty column F: not part of the list.
ws["G1"] = "Outside"
for r in range(2, 9):
    ws.cell(r, 7, f"g{r}")

wb.save(OUT)
print(OUT)
