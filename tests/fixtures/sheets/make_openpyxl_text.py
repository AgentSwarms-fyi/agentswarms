"""Text cells that read like numbers (R164).

Each value in column A is text in the file (a string cell), as Excel
keeps it: a price typed with its currency, a code that looks like a
number in scientific notation, a text that starts with an apostrophe.
Column B asks ISTEXT of each (Excel: TRUE), and C1:C2 count them
(Excel: COUNT 0, COUNTA 9), D1 sums them (0).
E1 is a number, as a control.

    python tests/fixtures/sheets/make_openpyxl_text.py
"""

import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "openpyxl-text.xlsx")

TEXTS = ["£1,234.50", "€99", "¥500", "1e5", "2E3", "'quoted", "$1,200", "007", "=1+1"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Texts"
for i, t in enumerate(TEXTS, start=1):
    c = ws.cell(i, 1, t)
    c.data_type = "s"  # "=1+1" too is text, not a formula
    ws.cell(i, 2, f"=ISTEXT(A{i})")
ws["C1"] = f"=COUNT(A1:A{len(TEXTS)})"
ws["C2"] = f"=COUNTA(A1:A{len(TEXTS)})"
ws["D1"] = f"=SUM(A1:A{len(TEXTS)})"
# A number stays a number.
ws["E1"] = 1234.5
ws["E2"] = "=ISNUMBER(E1)"
# A plain word, as a control: it comes in as typed.
ws["F1"] = "Hello"
wb.save(OUT)
print(OUT)
